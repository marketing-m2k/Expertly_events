"""One-off (re-runnable) pass over Master.xlsx's "Needs Review" sheet: asks
Gemini to read each event's own name/description text and decide its real
date, since these rows exist specifically because the resolved Date column
disagreed with what the event's own text says. Promotes genuinely-upcoming
events (with the corrected date) into "Verified Events"; drops anything
that turns out to already be in the past; leaves anything Gemini can't
confidently date alone in "Needs Review".
"""

import json
import os
from datetime import datetime

import openpyxl
from dotenv import load_dotenv
from google import genai
from google.genai import types

from scraper.excel_writer import format_sheet

load_dotenv()

MODEL = os.environ.get("GEMINI_MODEL", "gemini-3.1-flash-lite")
MASTER_PATH = "output/Master.xlsx"

SCHEMA = {
    "type": "object",
    "properties": {
        "corrected_date": {"type": "string", "description": "YYYY-MM-DD, or empty string if truly undeterminable"},
        "verdict": {"type": "string", "enum": ["upcoming", "past", "unclear"]},
    },
    "required": ["corrected_date", "verdict"],
}

PROMPT_TEMPLATE = """Today's date is {today}.

This event's resolved Date field ({resolved_date}) does not match what the
event's own name/description text states about when it actually happens --
that's why it's flagged for review instead of being trusted automatically.

Read the event's own text below and determine the REAL date this event
takes place. Prefer an explicit date/date-range stated in the text over the
resolved Date field if they disagree -- the resolved field is what's in
question here, not the text.

Respond with:
- corrected_date: the real date in YYYY-MM-DD format, if the text clearly
  states one (for a multi-day event, use the start date). Empty string if
  the text genuinely doesn't let you determine a reliable date.
- verdict: "upcoming" if that real date is on or after today, "past" if
  before today, "unclear" if you couldn't determine corrected_date at all.

Event name: {name}
Organizer: {organizer}
Description: {description}
"""


def resolve_all():
    client = genai.Client(api_key=os.environ["GEMINI_API_KEY"])

    wb = openpyxl.load_workbook(MASTER_PATH, data_only=True)
    ws_review = wb["Needs Review"]
    header = [c.value for c in ws_review[1]]
    rows = [list(r) for r in ws_review.iter_rows(min_row=2, values_only=True) if r and r[0]]

    today = datetime.now().strftime("%Y-%m-%d")
    today_dt = datetime.now().date()

    promoted, dropped, kept = [], [], []

    for row in rows:
        name, organizer, description = row[4], row[5], row[9] or ""
        prompt = PROMPT_TEMPLATE.format(
            today=today, resolved_date=row[1], name=name, organizer=organizer,
            description=description[:3000],
        )
        config = types.GenerateContentConfig(response_mime_type="application/json", response_schema=SCHEMA)
        response = client.models.generate_content(model=MODEL, contents=prompt, config=config)
        try:
            result = json.loads(response.text)
        except (ValueError, AttributeError, TypeError):
            kept.append(row)
            continue

        verdict = result.get("verdict")
        corrected = result.get("corrected_date") or ""

        if verdict == "upcoming" and corrected:
            try:
                corrected_dt = datetime.strptime(corrected, "%Y-%m-%d").date()
            except ValueError:
                kept.append(row)
                continue
            if corrected_dt < today_dt:
                # Gemini said "upcoming" but its own corrected date is in the
                # past -- contradictory, don't trust it either way.
                kept.append(row)
                continue
            new_row = list(row)
            new_row[1] = corrected_dt.strftime("%d-%b-%Y")  # Date
            promoted.append(new_row)
            print(f"PROMOTE: {name[:70]!r} -> {new_row[1]}")
        elif verdict == "past":
            dropped.append(row)
            print(f"DROP (past): {name[:70]!r} (Gemini's real date: {corrected or 'unstated'})")
        else:
            kept.append(row)
            print(f"KEEP in Needs Review (unclear): {name[:70]!r}")

    ws_verified = wb["Verified Events"]
    for row in promoted:
        ws_verified.append(row)
    format_sheet(ws_verified, header)

    wb.remove(ws_review)
    ws_new_review = wb.create_sheet("Needs Review")
    ws_new_review.append(header)
    for row in kept:
        ws_new_review.append(row)
    format_sheet(ws_new_review, header)

    # Keep the Stats sheet's "Final Master.xlsx" table honest -- it would
    # otherwise still show the pre-resolution Verified/Needs-Review counts.
    import weekly_full_run  # local import: avoids a circular import at module load time

    all_verified_rows = [list(r) for r in ws_verified.iter_rows(min_row=2, values_only=True) if r and r[0]]
    all_review_rows = [list(r) for r in ws_new_review.iter_rows(min_row=2, values_only=True) if r and r[0]]
    if "Stats" in wb.sheetnames:
        wb.remove(wb["Stats"])
    weekly_full_run._build_stats_sheet(wb, all_verified_rows, all_review_rows)
    wb.move_sheet("Stats", offset=-(len(wb.sheetnames) - 1))

    wb.save(MASTER_PATH)

    print(f"\nDone: {len(promoted)} promoted to Verified Events, {len(dropped)} dropped as past, "
          f"{len(kept)} left in Needs Review.")
    return len(promoted), len(dropped), len(kept)


if __name__ == "__main__":
    resolve_all()
