"""Stages 2-6 for one country: open every upcoming event's own page, extract
its details with proof, classify it, verify it, and re-scrape any event that
fails a check before giving up on it.

Reads the country's raw sheet (stage 1's list-page results) and returns one
record per event. Nothing here writes to Master; master_store.py does that.
"""

import hashlib
import html as html_lib
import os
from datetime import datetime

import openpyxl

from scraper.classify_rules import classify_event
from scraper.detail_extract import extract_detail_page
from scraper.event_id import make_event_id
from scraper.evidence import normalize_ws
from scraper.fetch import fetch_detail_pages
from scraper.label_extract import main_text, visible_text
from scraper.site_health import load_json, save_json
from scraper.verify_events import resolve_country, verify_event

MAX_RETRIES = 2


def read_upcoming_raw(path: str, sheet: str) -> list[dict]:
    """The country's raw rows that aren't already past. Only the identity
    columns are trusted from here (name, link, source); the list-page date,
    format and location came from the pattern matcher and are never used."""
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except FileNotFoundError:
        return []
    ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(min_row=1, values_only=True)
    header = [str(h or "").strip() for h in next(rows, [])]
    idx = {h: i for i, h in enumerate(header)}
    out = []
    for row in rows:
        if not row:
            continue
        get = lambda col: normalize_ws(str(row[idx[col]] or "")) if col in idx and idx[col] < len(row) else ""
        if get("Status") == "Past" or not get("Event Name"):
            continue
        out.append({"name": get("Event Name"), "organizer": get("Organizer"),
                    "link": get("Register Link"), "source_url": get("Source URL")})
    return out


def _haystack(html: str) -> str:
    """Everything a value might legitimately be found in: the page's source
    (structured data lives there) and its visible text."""
    return html_lib.unescape(html).replace("\\/", "/") + " " + visible_text(html)


def _fingerprint(html: str) -> str:
    return hashlib.sha1(main_text(html).encode("utf-8")).hexdigest()[:16]


def process_event(row: dict, fetched: dict, source_country: str, refetch, today: datetime, max_retries: int = MAX_RETRIES) -> dict:
    """Extract, classify and verify one event; on failure re-open its page up
    to `max_retries` more times (with a longer wait) before giving up."""
    link = row["link"]
    event_id = make_event_id(row["organizer"], row["name"], link, row["source_url"])
    base = {"event_id": event_id, "organizer": row["organizer"], "source_url": row["source_url"],
            "link": link, "source_country": source_country}

    attempts = 0
    result = fetched
    while True:
        if not result or result.get("html") is None:
            return {**base, "status": "unreachable", "fields": {}, "country": source_country, "category": "",
                    "reasons": [f"event page could not be opened ({(result or {}).get('error') or 'no result'})"],
                    "attempts": attempts, "fingerprint": ""}

        html = result["html"]
        extraction = extract_detail_page(html, link, row["organizer"])
        fields = extraction["fields"]
        name = fields.get("name", {}).get("value", "")
        description = fields.get("description", {}).get("value", "")
        classification = classify_event(name, description)
        failures = verify_event(fields, classification["category"], _haystack(html), today)
        failures += [r for r in extraction["review_reasons"]
                     if r not in failures and "no event date" not in r and " came from a " not in r]

        if classification["category"] == "Reject" or (not failures and not extraction["conflicts"]) or attempts >= max_retries:
            break
        attempts += 1
        result = refetch(link, wait_ms=3000 * attempts)

    if classification["category"] == "Reject":
        status = "rejected"
        reasons = [classification["reason"]]
    elif failures or extraction["conflicts"]:
        status = "needs_review"
        reasons = failures
    else:
        status = "verified"
        reasons = []

    location = fields.get("location", {}).get("value", "")
    return {**base, "status": status, "fields": fields, "category": classification["category"],
            "classification_reason": classification["reason"], "reasons": reasons, "attempts": attempts,
            "country": resolve_country(location, source_country), "fingerprint": _fingerprint(html)}


def enrich_country(label: str, raw_path: str, raw_sheet: str, state_dir: str, today: datetime | None = None,
                   fetch_fn=fetch_detail_pages, workers: int = 5, limit: int = 0) -> list[dict]:
    today = today or datetime.now()
    rows = read_upcoming_raw(raw_path, raw_sheet)
    if limit:
        rows = rows[:limit]

    state_path = os.path.join(state_dir, f"pages_{label}.json")
    state = load_json(state_path)

    seen, unique_rows = set(), []
    for r in rows:  # the same event can appear twice in the raw sheet
        eid = make_event_id(r["organizer"], r["name"], r["link"], r["source_url"])
        if eid not in seen:
            seen.add(eid)
            unique_rows.append(r)

    fetched = fetch_fn([r["link"] for r in unique_rows if r["link"]], workers=workers)
    refetch = lambda url, wait_ms: fetch_fn([url], wait_ms=wait_ms, workers=1).get(url)

    records = []
    for r in unique_rows:
        if not r["link"]:
            eid = make_event_id(r["organizer"], r["name"], r["link"], r["source_url"])
            records.append({"event_id": eid, "organizer": r["organizer"], "source_url": r["source_url"], "link": "",
                            "source_country": label, "status": "needs_review", "fields": {}, "country": label,
                            "category": "", "reasons": ["the list page gave no link to the event's own page"],
                            "attempts": 0, "fingerprint": "", "changed": True})
            continue

        result = fetched.get(r["link"])
        record = process_event(r, result, label, refetch, today)
        previous = state.get(record["event_id"])
        page_ok = record["status"] != "unreachable"
        if page_ok and previous and previous.get("fingerprint") == record["fingerprint"] and previous.get("record"):
            # page content unchanged since last week: keep last week's values
            # rather than re-deriving them (and, later, re-paying for review)
            record = {**previous["record"], "changed": False}
        elif not page_ok and previous and previous.get("record"):
            record = {**previous["record"], "changed": False, "unreachable_this_run": True}
        else:
            record["changed"] = True
        if page_ok:
            state[record["event_id"]] = {"fingerprint": record["fingerprint"], "record": {k: v for k, v in record.items() if k != "changed"}}
        records.append(record)

    save_json(state_path, state)
    return records
