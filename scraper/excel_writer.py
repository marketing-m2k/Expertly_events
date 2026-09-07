"""Append scraped events to an Excel workbook, skipping duplicates."""

import re
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from scraper.date_utils import has_day_precision, is_before_cutoff, is_past_event, parse_event_date

# Only events on/after this date are kept — the user wants Jan 2026 onward,
# past and upcoming both, not the full historical archive of every site.
CUTOFF_DATE = datetime(2026, 1, 1)

# Reading order tuned for a human scanning the sheet: when/what/who first,
# the long free-text description last, links at the very end. No "Country"
# column — every site scraped is India-only, so it was 100% blank dead
# weight in every row.
COLUMNS = [
    "Date",
    "Status",
    "Event Name",
    "Organizer",
    "Category",
    "Format",
    "Location",
    "Description",
    "Register Link",
    "Source URL",
]

# Used for the final cleaned/classified workbooks (clean_events.py) only --
# the raw scraped workbooks (COLUMNS above) keep whatever single date string
# was scraped as-is; the End Date split into its own column happens during
# cleaning, once a date range like "1st to 18th September, 2026" has been
# resolved into a real start/end pair.
FINAL_COLUMNS = [
    "Date",
    "End Date",
    "Status",
    "Event Name",
    "Organizer",
    "Category",
    "Format",
    "Location",
    "Description",
    "Register Link",
    "Source URL",
]

_COLUMN_WIDTHS = {
    "Date": 20,
    "End Date": 20,
    "Status": 11,
    "Event Name": 48,
    "Organizer": 34,
    "Category": 26,
    "Format": 11,
    "Location": 16,
    "Description": 55,
    "Register Link": 40,
    "Source URL": 40,
}

_HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_WRAP_COLS = {"Event Name", "Description"}


def format_sheet(ws, columns=None):
    """Apply consistent, readable formatting to a sheet already holding
    columns-shaped rows (defaults to COLUMNS): sized columns, wrapped
    long-text columns, a bold header row frozen in place, the whole range
    turned into a filterable/banded Excel Table, and clickable hyperlinks
    on the two link columns."""
    columns = columns or COLUMNS
    if ws.max_row < 1:
        return

    for col_idx, name in enumerate(columns, start=1):
        header_cell = ws.cell(row=1, column=col_idx, value=name)
        header_cell.fill = _HEADER_FILL
        header_cell.font = _HEADER_FONT
        header_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = _COLUMN_WIDTHS.get(name, 20)

    wrap_alignment = Alignment(wrap_text=True, vertical="top")
    plain_alignment = Alignment(vertical="top")
    link_col_idx = {name: i + 1 for i, name in enumerate(columns)}
    for row_idx in range(2, ws.max_row + 1):
        for col_idx, name in enumerate(columns, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = wrap_alignment if name in _WRAP_COLS else plain_alignment
        for col_name in ("Register Link", "Source URL"):
            cell = ws.cell(row=row_idx, column=link_col_idx[col_name])
            url = cell.value
            if url and str(url).startswith("http"):
                cell.hyperlink = url
                cell.style = "Hyperlink"

    ws.row_dimensions[1].height = 20

    # remove any pre-existing table definition before re-adding — openpyxl
    # errors on a duplicate table name if one is already on the sheet from a
    # prior formatting pass
    for name in list(ws.tables.keys()):
        del ws.tables[name]

    last_col_letter = get_column_letter(len(columns))
    table_ref = f"A1:{last_col_letter}{ws.max_row}"
    safe_name = "Tbl_" + re.sub(r"\W+", "_", ws.title).strip("_")
    table = Table(displayName=safe_name, ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9", showRowStripes=True, showFirstColumn=False,
        showLastColumn=False, showColumnStripes=False,
    )
    ws.add_table(table)


def _dedupe_key(row: dict) -> tuple:
    # Organizer is deliberately excluded: the same event often gets pulled
    # from two different sub-pages of the same institution (e.g. ICAI's CPE
    # Directorate page and its Research committee page both mention the same
    # webinar), each tagging it with a different organizer label. Keying on
    # name+date only collapses those into one row instead of duplicating it.
    return (
        (row.get("event_name") or "").strip().lower(),
        (row.get("date") or "").strip().lower(),
    )


def upsert_events(path: str, sheet_name: str, events: list[dict]) -> tuple[int, int]:
    """Add new events and, for events already on the sheet with an
    incomplete date (blank or day-less, e.g. "APR"), fix the existing row in
    place if this scrape found a precise date for the same event name —
    instead of adding a second row for what a site simply updated. Returns
    (added, updated)."""
    try:
        wb = openpyxl.load_workbook(path)
    except FileNotFoundError:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)
        ws.append(COLUMNS)

    existing_keys = set()
    # event name -> row index (1-based), for existing rows whose stored date
    # has no day-level precision -- candidates to be repaired in place.
    incomplete_by_name: dict[str, int] = {}
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row:
            continue
        name_norm = str(row[2] or "").strip().lower()
        date_raw = row[0] or ""
        existing_keys.add((name_norm, str(date_raw).strip().lower()))
        if name_norm and not has_day_precision(date_raw):
            incomplete_by_name[name_norm] = row_idx

    added = 0
    updated = 0
    update_fields = (
        (4, "organizer"), (5, "category"), (6, "format"),
        (7, "location_city"), (8, "description"), (9, "link"), (10, "source_url"),
    )
    for event in events:
        date = event.get("date", "")
        if is_before_cutoff(date, CUTOFF_DATE):
            continue
        name_norm = (event.get("event_name") or "").strip().lower()
        key = (name_norm, str(date).strip().lower())
        if key in existing_keys:
            continue

        status = "Past" if is_past_event(date) else "Upcoming"

        if has_day_precision(date) and name_norm in incomplete_by_name:
            row_idx = incomplete_by_name.pop(name_norm)
            ws.cell(row=row_idx, column=1, value=date)
            ws.cell(row=row_idx, column=2, value=status)
            for col_idx, field in update_fields:
                cell = ws.cell(row=row_idx, column=col_idx)
                new_val = event.get(field, "")
                if not cell.value and new_val:
                    cell.value = new_val
            existing_keys.add(key)
            updated += 1
            continue

        existing_keys.add(key)
        ws.append([
            date,
            status,
            event.get("event_name", ""),
            event.get("organizer", ""),
            event.get("category", ""),
            event.get("format", ""),
            event.get("location_city", ""),
            event.get("description", ""),
            event.get("link", ""),
            event.get("source_url", ""),
        ])
        added += 1

    wb.save(path)
    return added, updated


def sort_and_link(path: str, sheet_name: str):
    """Sort the sheet chronologically (oldest-first, so Jan 2026 past events
    lead into upcoming ones), refresh each row's Status, drop rows older
    than CUTOFF_DATE, and turn URL columns into clickable hyperlinks."""
    wb = openpyxl.load_workbook(path)
    if sheet_name not in wb.sheetnames:
        return
    ws = wb[sheet_name]

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    if not rows:
        return

    today = datetime.now()
    rows = [r for r in rows if not is_before_cutoff(r[0], CUTOFF_DATE, today)]
    rows = [list(r) + [""] * (len(COLUMNS) - len(r)) for r in rows]
    for r in rows:
        r[1] = "Past" if is_past_event(r[0], today) else "Upcoming"

    def sort_key(row):
        dt = parse_event_date(row[0], today)
        return dt if dt is not None else datetime.max

    rows.sort(key=sort_key)

    ws.delete_rows(2, ws.max_row)
    for row in rows:
        ws.append(row)

    format_sheet(ws)
    wb.save(path)
