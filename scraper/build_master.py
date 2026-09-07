"""Build Master.xlsx: one "Verified Events" sheet listing every verified
event across every country's final workbook.

Also captures any row a human has manually marked "Verified" in a
country's Upcoming - Incomplete sheet (only if the Date cell has actually
been filled in with a real date -- typing "Verified" alone isn't enough)
into scraper/manual_overrides.py's persistent override file. That file is
what makes the correction stick: clean_events.py checks it on every future
weekly run, so a manually-verified event lands back in that country's own
Upcoming - Verified sheet automatically instead of reverting to Incomplete.

Run this whenever you've finished marking some rows Verified and want
those changes reflected here:

    python -m scraper.build_master
"""

import openpyxl

from scraper import manual_overrides
from scraper.date_utils import has_day_precision
from scraper.excel_writer import FINAL_COLUMNS, format_sheet

# Add a country here (label -> its final workbook path) and it's picked up
# automatically -- no other change needed for build_master.py itself.
COUNTRY_FILES = {
    "India": "output/Events_2026.xlsx",
    "USA": "output/Events_USA_2026.xlsx",
}

MASTER_PATH = "output/Master.xlsx"
MASTER_COLUMNS = ["Country"] + FINAL_COLUMNS


def _capture_manual_verifications(label: str, path: str, overrides: dict) -> int:
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except FileNotFoundError:
        print(f"  [{label}] {path} not found -- skipping")
        return 0
    if "Upcoming - Incomplete" not in wb.sheetnames:
        return 0

    ws = wb["Upcoming - Incomplete"]
    header = [c.value for c in ws[1]]
    idx = {name: i for i, name in enumerate(header)}
    captured = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        status = str(row[idx.get("Status", -1)] or "").strip().lower()
        if status != "verified":
            continue
        name = row[idx.get("Event Name", -1)]
        date_val = row[idx.get("Date", -1)]
        if not has_day_precision(str(date_val or "")):
            print(f"  [{label}] '{name}' is marked Verified but Date isn't a real filled-in "
                  f"date -- fill in Date (and End Date if multi-day) before marking Verified. Skipped.")
            continue
        manual_overrides.record(overrides, label, name, list(row))
        captured += 1

    return captured


def build() -> None:
    overrides = manual_overrides.load()
    total_captured = 0
    for label, path in COUNTRY_FILES.items():
        total_captured += _capture_manual_verifications(label, path, overrides)
    manual_overrides.save(overrides)
    print(f"Captured {total_captured} newly manually-verified event(s) "
          f"({len(overrides)} total on record) into {manual_overrides.PATH}")

    all_rows = []
    for label, path in COUNTRY_FILES.items():
        try:
            src = openpyxl.load_workbook(path, data_only=True)
        except FileNotFoundError:
            continue
        ws_v = src["Upcoming - Verified"]
        for row in ws_v.iter_rows(min_row=2, values_only=True):
            if row and (row[3] or "").strip():  # Event Name non-blank
                all_rows.append([label] + list(row))

    # An override captured just now (or in an earlier run) won't show up in
    # this week's Verified sheet until clean_events.py re-processes it next
    # run -- include it in Master immediately rather than waiting a week.
    seen = {(r[0], (r[4] or "").strip().lower()) for r in all_rows}
    for rec in overrides.values():
        pair = (rec["country"], (rec["name"] or "").strip().lower())
        if pair not in seen:
            all_rows.append([rec["country"]] + rec["row"])
            seen.add(pair)

    def sort_key(row):
        return (row[0], row[1] or "")  # Country, then Date string

    all_rows.sort(key=sort_key)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Verified Events")
    ws.append(MASTER_COLUMNS)
    for r in all_rows:
        ws.append(r)
    format_sheet(ws, MASTER_COLUMNS)
    wb.save(MASTER_PATH)
    print(f"Wrote {MASTER_PATH}: {len(all_rows)} verified events across {len(COUNTRY_FILES)} countries.")


if __name__ == "__main__":
    build()
