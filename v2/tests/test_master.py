"""Phase 9-11: Master / de-duplication / history / locks / versioned records."""

import json
import os

import openpyxl
import pytest

from pipeline import identity
from pipeline.master import export_master, import_locks
from pipeline.models import Status
from pipeline.store import Store
from tests.pages import TODAY


def rec(eid, name="Tax Update", url=None, date="2026-11-10", location="", status=Status.VERIFIED, org="Example Institute", **kw):
    url = url if url is not None else f"https://inst.org/events/{eid}"
    keys = identity.identity_keys(org, url, kw.pop("shared", False), name, date, location, kw.pop("ext_id", ""))
    fields = {"name": {"value": name, "source": "heading"}, "start_date": {"value": date, "source": "label"}, "venue": {"value": location, "source": "label"} if location else {"value": None}}
    base = {"event_id": eid, "org": org, "country": "UK", "source_url": "https://inst.org/events", "url": url, "status": status, "category": "Tax", "reasons": [],
            "flags": [], "fields": fields, "packet": {}, "extractor_version": "2.0.0", **keys, "verified_by": "rules"}
    base.update(kw)
    return base


@pytest.fixture
def st(tmp_path):
    return Store(str(tmp_path / "s.sqlite"))


def test_same_name_and_date_in_different_locations_are_two_events(st):
    a = st.upsert(rec("ev_a", "Tax Update", url="", location="Chicago", shared=True), "r1")
    b = st.upsert(rec("ev_b", "Tax Update", url="", location="Dallas", shared=True), "r1")
    assert a[0] == "added" and b[0] == "added" and a[1] != b[1]
    assert len(st.rows("dup_of IS NULL")) == 2


def test_same_url_is_the_same_event_even_when_the_name_changes(st):
    st.upsert(rec("ev_a", "GST Workshop", url="https://inst.org/events/gst"), "r1")
    action, eid = st.upsert(rec("ev_zzz", "GST Workshop 2026 (renamed)", url="https://www.inst.org/events/gst/"), "r2")
    assert action == "updated" and eid == "ev_a"
    assert st.get("ev_a")["name"].endswith("(renamed)")
    assert len(st.rows()) == 1


def test_same_name_date_location_on_different_urls_is_merged_and_the_other_url_is_kept(st):
    st.upsert(rec("ev_a", "GST Workshop", url="https://inst.org/events/gst", location="Hilton Chennai"), "r1")
    action, eid = st.upsert(rec("ev_b", "GST Workshop", url="https://inst.org/calendar/gst-copy", location="Hilton Chennai"), "r1")
    assert action == "duplicate_of" and eid == "ev_a" and st.get("ev_b") is None
    assert json.loads(st.get("ev_a")["alt_urls_json"]) == ["https://inst.org/calendar/gst-copy"]
    assert any("also listed at" in h["reason"] for h in st.history("status_history", "ev_a"))


def test_missing_location_on_one_side_is_flagged_possible_duplicate_never_merged(st):
    st.upsert(rec("ev_a", "GST Workshop", url="https://inst.org/events/gst", location="Hilton Chennai"), "r1")
    action, eid = st.upsert(rec("ev_b", "GST Workshop", url="https://inst.org/events/gst-b", location=""), "r1")
    assert action == "added" and st.get(eid)["dup_of"] is None
    assert "ev_a" in (st.get(eid)["possible_dup_of"] or "")


def test_status_history_records_every_transition_with_versions(st):
    st.upsert(rec("ev_a", status=Status.QUALITY_REVIEW, reasons=["missing_date"]), "r1")
    st.upsert(rec("ev_a", status=Status.NEEDS_REVIEW, reasons=["still unclear"]), "r2")
    st.upsert(rec("ev_a", status=Status.VERIFIED, reasons=["ok"]), "r3")
    h = st.history("status_history", "ev_a")
    assert [(x["from_status"], x["to_status"]) for x in h] == [(None, "QUALITY_REVIEW"), ("QUALITY_REVIEW", "NEEDS_REVIEW"), ("NEEDS_REVIEW", "VERIFIED")]
    assert h[0]["pipeline_version"] and h[0]["extractor_version"] and h[0]["run_id"] == "r1"


def test_history_is_kept_first_discovered_never_overwritten_and_changed_fields_logged(st):
    st.upsert(rec("ev_a", date="2026-11-10"), "r1")
    first = st.get("ev_a")["first_discovered"]
    st.upsert(rec("ev_a", date="2026-11-12"), "r2")
    row = st.get("ev_a")
    assert row["first_discovered"] == first and row["start_date"] == "2026-11-12"
    fh = st.history("field_history", "ev_a")
    assert any(x["field"] == "start_date" and x["old_value"] == "2026-11-10" and x["new_value"] == "2026-11-12" for x in fh)


def test_a_worse_run_never_overwrites_a_verified_record(st):
    st.upsert(rec("ev_a", status=Status.VERIFIED), "r1")
    action, _ = st.upsert(rec("ev_a", status=Status.FETCH_FAILED, reasons=["HTTP 503"], fields={"name": {"value": None}}), "r2")
    row = st.get("ev_a")
    assert action == "kept_previous" and row["status"] == "VERIFIED" and row["name"] == "Tax Update"
    assert "FETCH_FAILED" in (row["last_attempt_json"] or "")
    assert any("could not re-verify" in x["reason"] for x in st.history("status_history", "ev_a"))


def test_unchanged_record_is_left_alone(st):
    r = rec("ev_a")
    st.upsert(r, "r1")
    assert st.upsert(dict(r), "r2")[0] == "unchanged"


def test_locked_rows_are_never_touched(st):
    st.upsert(rec("ev_a", "Original Name"), "r1")
    st.set_lock("ev_a", True)
    action, _ = st.upsert(rec("ev_a", "Scraper Wants To Change This"), "r2")
    assert action == "locked_skip" and st.get("ev_a")["name"] == "Original Name"


def test_ai_corrections_are_kept_with_their_evidence(st):
    st.upsert(rec("ev_a"), "r1")
    st.log_correction("ev_a", "r1", {"field": "venue", "old": "", "new": "Hilton", "evidence_quote": "Venue: Hilton", "source_url": "https://inst.org/e", "reason": "stated"})
    h = st.history("correction_history", "ev_a")
    assert h[0]["evidence_quote"] == "Venue: Hilton" and h[0]["by"] == "gemini"


def test_rejected_records_keep_reason_evidence_url_timestamp_and_versions(st):
    rej = {"reason": "a social golf day", "evidence_quote": "Join us for golf", "source_url": "https://inst.org/e"}
    st.upsert(rec("ev_a", status=Status.REJECTED, rejection=rej, reasons=["AI review"]), "r1")
    row = st.get("ev_a")
    assert json.loads(row["rejection_json"]) == rej and row["verified_at"] and row["pipeline_version"] and row["extractor_version"]


def test_a_rejected_event_is_eligible_again_when_the_extractor_changes(st):
    from pipeline.reprocess import select_for_reprocess
    st.upsert(rec("ev_a", status=Status.REJECTED, extractor_version="1.0.0"), "r1")
    assert [a.event_id for a in select_for_reprocess(st)] == ["ev_a"]


def test_master_contains_only_verified_events_and_keeps_everything_else_visible(st, tmp_path):
    st.upsert(rec("ev_v", "Verified One", status=Status.VERIFIED), "r1")
    st.upsert(rec("ev_c", "Corrected One", status=Status.CORRECTED), "r1")
    st.upsert(rec("ev_n", "Needs Review One", status=Status.NEEDS_REVIEW, reasons=["x"]), "r1")
    st.upsert(rec("ev_q", "Quality One", status=Status.QUALITY_REVIEW), "r1")
    st.upsert(rec("ev_r", "Rejected One", status=Status.REJECTED, rejection={"reason": "r", "evidence_quote": "q", "source_url": "u"}), "r1")
    st.upsert(rec("ev_f", "Failed One", status=Status.FETCH_FAILED, reasons=["HTTP 404"]), "r1")
    st.upsert(rec("ev_p", "Past One", status=Status.PAST, date="2026-01-10"), "r1")
    path = str(tmp_path / "Master_v2.xlsx")
    export_master(st, path, TODAY)
    wb = openpyxl.load_workbook(path)
    def names(sheet):
        ws = wb[sheet]
        hdr = [c.value for c in ws[1]]
        i = hdr.index("Event Name")
        return {r[i] for r in ws.iter_rows(min_row=2, values_only=True) if r[i]}
    assert names("Verified Events") == {"Verified One", "Corrected One"}
    assert names("Needs Review") == {"Needs Review One", "Quality One"}
    assert names("Rejected") == {"Rejected One"} and names("Fetch Failed") == {"Failed One"} and names("Past Events") == {"Past One"}
    rej = wb["Rejected"]
    hdr = [c.value for c in rej[1]]
    row = [c.value for c in rej[2]]
    assert row[hdr.index("Evidence Quote")] == "q" and row[hdr.index("Official URL")] == "u"
    assert {"Status History", "Correction History", "Org Health", "Duplicates"} <= set(wb.sheetnames)


def test_master_keeps_two_same_named_events_in_different_places(st, tmp_path):
    st.upsert(rec("ev_a", "Tax Update", url="", location="Chicago", shared=True), "r1")
    st.upsert(rec("ev_b", "Tax Update", url="", location="Dallas", shared=True), "r1")
    path = str(tmp_path / "M.xlsx")
    export_master(st, path, TODAY)
    ws = openpyxl.load_workbook(path)["Verified Events"]
    hdr = [c.value for c in ws[1]]
    locs = sorted(r[hdr.index("Location")] for r in ws.iter_rows(min_row=2, values_only=True) if r[0])
    assert locs == ["Chicago", "Dallas"]


def test_locks_survive_a_master_round_trip(st, tmp_path):
    st.upsert(rec("ev_a", "Locked Event"), "r1")
    path = str(tmp_path / "M.xlsx")
    export_master(st, path, TODAY)
    wb = openpyxl.load_workbook(path)
    ws = wb["Verified Events"]
    hdr = [c.value for c in ws[1]]
    ws.cell(row=2, column=hdr.index("Locked") + 1, value="Yes")        # a person locks the row in Excel
    wb.save(path)
    assert import_locks(st, path) == 1 and st.get("ev_a")["locked"] == 1


def test_event_ids_follow_the_hierarchy():
    a = identity.make_event_id("O", "https://www.inst.org/e/1/", shared_url=False, name="A", start_date="2026-11-10", location="x")
    b = identity.make_event_id("O", "https://inst.org/e/1?utm_source=x", shared_url=False, name="Renamed", start_date="2026-11-11", location="y")
    assert a == b                                                         # URL identity beats name/date/location
    c = identity.make_event_id("O", "https://inst.org/series", shared_url=True, name="Tax Update", start_date="2026-11-10", location="Chicago")
    d = identity.make_event_id("O", "https://inst.org/series", shared_url=True, name="Tax Update", start_date="2026-11-10", location="Dallas")
    assert c != d                                                         # a shared URL is not an identity
    e = identity.make_event_id("O", "", shared_url=True, name="X", start_date="2026-11-10", location="", ext_id="UID-1")
    f = identity.make_event_id("O", "", shared_url=True, name="Different", start_date="2026-12-10", location="", ext_id="UID-1")
    assert e == f                                                         # external id, same org


def test_raw_pages_are_preserved_compressed(st):
    h = st.save_page("ev_a", "https://inst.org/e", "<html>hello evidence</html>", 200, "http")
    assert h and st.load_page("ev_a") == "<html>hello evidence</html>"
