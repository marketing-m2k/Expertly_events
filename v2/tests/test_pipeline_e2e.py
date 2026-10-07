"""End to end on a tiny fake internet: every discovered event must be traceable; scraper failures are never rejections;
incomplete organisations never cause 'not found' demotions; the stages run in order and the QC report accounts for everything."""

import json
import os
from datetime import date

import openpyxl
import pytest

from pipeline.fetcher import Page, ReplayFetcher
from pipeline.models import Status
from pipeline.qc import audit_ledger
from pipeline.run import RunConfig, run_pipeline
from pipeline.store import Store
from tests.fakes import FakeLLM, verified
from tests.helpers import card_article, page
from tests.pages import TODAY, event_page, jsonld_event

BASE = "https://inst.org"
LIST_URL = BASE + "/events"
DESC = "<p>" + "A practical session for tax practitioners and finance teams on current developments, with worked examples and time for questions. " * 2 + "</p>"


def listing():
    cards = "".join([
        f"<article><h3><a href='{BASE}/events/gst-audit-workshop'>GST Audit Workshop</a></h3><p>12 November 2026</p></article>",
        f"<article><h3><a href='{BASE}/events/webinar-777'>Webinar</a></h3><p>13 November 2026</p></article>",
        f"<article><h3><a href='{BASE}/events/missing-page-seminar'>Transfer Pricing Seminar</a></h3><p>14 November 2026</p></article>",
        f"<article><h3><a href='{BASE}/events/js-only-forum'>Insolvency Forum</a></h3><p>15 November 2026</p></article>",
        f"<article><h3><a href='{BASE}/events/old-tax-day'>Old Tax Day</a></h3><p>1 March 2026</p></article>"])
    return page("<h1>Events</h1>" + cards)


def site():
    return {
        LIST_URL: listing(),
        BASE + "/events/gst-audit-workshop": event_page(date_line="12 November 2026", fmt="In person", address="2 Mount Road, Chennai", desc=DESC,
                                                         jsonld=None),
        BASE + "/events/webinar-777": event_page(title="Indirect Tax Webinar", h1="Indirect Tax Webinar", date_line="13 November 2026", fmt="Online", venue="Zoom", desc=DESC),
        # /events/missing-page-seminar -> 404
        BASE + "/events/js-only-forum": "<html><body><div id='root'></div><script>x()</script></body></html>",
        BASE + "/events/old-tax-day": event_page(title="Old Tax Day", h1="Old Tax Day", date_line="1 March 2026", fmt="In person", desc=DESC),
    }


@pytest.fixture
def out(tmp_path):
    return str(tmp_path / "run")


def go(out, pages=None, llm=None, orgs=None, review="all", store=None):
    pages = pages or site()
    orgs = orgs or [{"organizer": "Example Institute", "country": "UK", "url": LIST_URL}]
    cfg = RunConfig(out_dir=out, review_mode=review, workers=2, discovery_workers=1, today=TODAY, use_browser=False, use_feeds=False)
    return run_pipeline(orgs, cfg, llm=llm or FakeLLM(default=verified()), http=ReplayFetcher(pages), store=store)


def statuses(result):
    return {r["name"]: r["status"] for r in result["store"].rows()}


def test_every_discovered_event_gets_a_final_outcome(out):
    r = go(out)
    audit = r["qc"]["audit"]
    assert audit["discovered_candidates"] == 5 and audit["with_final_status"] == 5 and audit["lost"] == []
    assert r["qc"]["headline"]["TRACEABILITY: candidates without a final outcome"] == 0


def test_fetch_failure_is_FETCH_FAILED_never_rejected(out):
    r = go(out)
    rows = {x["url"]: x for x in r["store"].rows()}
    failed = rows[BASE + "/events/missing-page-seminar"]
    assert failed["status"] == Status.FETCH_FAILED and "404" in failed["reasons_json"]
    assert Status.REJECTED not in {x["status"] for x in rows.values()}


def test_unreadable_page_is_held_for_review_not_rejected(out):
    r = go(out)
    row = {x["url"]: x for x in r["store"].rows()}[BASE + "/events/js-only-forum"]
    assert row["status"] in (Status.QUALITY_REVIEW, Status.NEEDS_REVIEW) and "page_unreadable" in row["flags_json"]


def test_statuses_after_a_full_run(out):
    r = go(out)
    st = {x["url"].rsplit("/", 1)[-1]: x["status"] for x in r["store"].rows()}
    assert st["gst-audit-workshop"] == Status.VERIFIED
    assert st["webinar-777"] == Status.VERIFIED                         # the generic card title was replaced by the event page's real name
    assert st["old-tax-day"] == Status.PAST
    name = {x["url"].rsplit("/", 1)[-1]: x["name"] for x in r["store"].rows()}
    assert name["webinar-777"] == "Indirect Tax Webinar"


def test_stages_leave_a_trail_in_the_ledger(out):
    r = go(out)
    store = r["store"]
    with store.lock:
        stages = {x["stage"] for x in store.db.execute("SELECT DISTINCT stage FROM ledger").fetchall()}
    assert {"discovery", "fetch", "extract", "quality", "verify", "classify", "review", "final", "master"} <= stages


def test_qc_report_files_and_headline_numbers(out):
    r = go(out)
    h = r["qc"]["headline"]
    assert h["Organisations processed"] == 1 and h["  discovery complete"] == 1
    assert h["Events discovered (candidates)"] == 5 and h["FETCH_FAILED"] == 1 and h["PAST"] == 1 and h["REJECTED"] == 0
    assert h["Gemini: events reviewed"] >= 2
    assert os.path.exists(os.path.join(out, "qc_report.md")) and os.path.exists(os.path.join(out, "qc_report.json"))
    md = open(os.path.join(out, "qc_report.md"), encoding="utf-8").read()
    assert "lost (no outcome): 0" in md


def test_master_only_has_verified_events_on_the_main_sheet(out):
    r = go(out)
    ws = openpyxl.load_workbook(r["master"])["Verified Events"]
    hdr = [c.value for c in ws[1]]
    names = {row[hdr.index("Event Name")] for row in ws.iter_rows(min_row=2, values_only=True) if row[0]}
    assert names == {"GST Audit Workshop", "Indirect Tax Webinar"}


def test_failed_and_incomplete_discovery_are_never_empty_and_never_demote(out):
    store = Store(os.path.join(out, "store.sqlite"))
    # a previously verified event of an organisation whose list page now fails
    old = {"event_id": "ev_old", "org": "Broken Org", "country": "UK", "source_url": "https://broken.org/events", "url": "https://broken.org/events/x",
           "status": Status.VERIFIED, "category": "Tax", "reasons": [], "flags": [], "packet": {}, "extractor_version": "2.0.0", "url_key": "broken.org/events/x",
           "fields": {"name": {"value": "Old Verified Event"}, "start_date": {"value": "2026-12-01"}}}
    store.upsert(old, "run_prev")
    with store.lock:
        store.db.execute("UPDATE events SET last_discovered='2026-08-01T00:00:00Z' WHERE event_id='ev_old'")
        store.db.commit()
    pages = site()
    pages["https://broken.org/events"] = Page(url="https://broken.org/events", status=500, error="HTTP 500")
    orgs = [{"organizer": "Example Institute", "country": "UK", "url": LIST_URL}, {"organizer": "Broken Org", "country": "UK", "url": "https://broken.org/events"}]
    r = go(out, pages=pages, orgs=orgs, store=store)
    states = {x.org: x.completeness["state"] for x in r["discovery"]}
    assert states["Broken Org"] == "discovery_failed" and states["Example Institute"] == "discovery_complete"
    assert store.get("ev_old")["status"] == Status.VERIFIED              # a failed organisation is never evidence that its events vanished
    assert r["qc"]["problem_orgs"][0]["org"] == "Broken Org"


def test_a_completely_read_organisation_that_stopped_listing_an_event_sends_it_to_review(out):
    store = Store(os.path.join(out, "store.sqlite"))
    old = {"event_id": "ev_gone", "org": "Example Institute", "country": "UK", "source_url": LIST_URL, "url": BASE + "/events/vanished",
           "status": Status.VERIFIED, "category": "Tax", "reasons": [], "flags": [], "packet": {}, "extractor_version": "2.0.0", "url_key": "inst.org/events/vanished",
           "fields": {"name": {"value": "Vanished Event"}, "start_date": {"value": "2026-12-01"}}}
    store.upsert(old, "run_prev")
    with store.lock:
        store.db.execute("UPDATE events SET last_discovered='2026-08-01T00:00:00Z' WHERE event_id='ev_gone'")
        store.db.commit()
    r = go(out, store=store)
    assert store.get("ev_gone")["status"] == Status.NEEDS_REVIEW
    assert any("not found on its website" in h["reason"] for h in store.history("status_history", "ev_gone"))


def test_rerun_is_idempotent_and_keeps_history(out):
    r1 = go(out)
    store = r1["store"]
    n1 = len(store.rows())
    r2 = go(out, store=store)
    assert len(store.rows()) == n1
    up = r2["stats"].upsert
    assert up.get("added", 0) == 0 and up.get("unchanged", 0) + up.get("updated", 0) + up.get("kept_previous", 0) >= 4


def test_ai_off_never_marks_anything_verified_that_needs_judgement(out):
    r = go(out, review="off", llm=FakeLLM())
    st = statuses(r)
    assert Status.REJECTED not in st.values()
    assert st["GST Audit Workshop"] == Status.VERIFIED


def test_gemini_budget_exhaustion_keeps_events_awaiting_review_not_rejected(out):
    from pipeline.llm import BudgetExceeded
    r = go(out, llm=FakeLLM(default=BudgetExceeded("cap")))
    st = set(statuses(r).values())
    assert Status.NEEDS_AI_REVIEW in st and Status.REJECTED not in st
    assert r["qc"]["headline"]["Gemini: unavailable / skipped"].startswith(("3", "2", "4"))


def test_the_cli_refuses_a_full_run_without_the_approval_phrase(capsys):
    import run_v2
    rc = run_v2.main(["--countries", "India", "USA", "UK", "SG", "UAE", "AUS", "--out", "data/never"])
    assert rc == 3 and "RUN FULL" in capsys.readouterr().out
